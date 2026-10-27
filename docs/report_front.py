"""The report's front matter: title page through symbols and notations.

Separated from make_report.py only to keep each file a readable length; it is
called from there and uses its helpers.
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

CENTER = WD_ALIGN_PARAGRAPH.CENTER
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT

ABBREVIATIONS = [
    ("API", "Application Programming Interface"),
    ("ASR", "Automatic Speech Recognition"),
    ("CLIP", "Contrastive Language-Image Pre-training"),
    ("CPU", "Central Processing Unit"),
    ("DFD", "Data-Flow Diagram"),
    ("EMA", "Exponential Moving Average"),
    ("FAISS", "Facebook AI Similarity Search"),
    ("GPU", "Graphics Processing Unit"),
    ("HMAC", "Hash-based Message Authentication Code"),
    ("JWT", "JSON Web Token"),
    ("KTS", "Kernel Temporal Segmentation"),
    ("LLM", "Large Language Model"),
    ("OCR", "Optical Character Recognition"),
    ("OTP", "One-Time Password"),
    ("REST", "Representational State Transfer"),
    ("SLA", "Service-Level Agreement"),
    ("SLO", "Service-Level Objective"),
    ("SQL", "Structured Query Language"),
    ("TRL", "Technology Readiness Level"),
    ("UML", "Unified Modelling Language"),
    ("VIT", "Vellore Institute of Technology"),
]

SYMBOLS = [
    ("c(x)", "predicted execution time for an input of x minutes, in seconds"),
    ("a", "fixed-overhead term of the execution-cost model, in seconds"),
    ("b", "per-input-minute rate of the execution-cost model"),
    ("n, Σx, Σy, Σxx, Σxy",
     "the five running sums from which the affine fit is recovered"),
    ("T", "the wall-clock budget, or deadline, allowed for a job"),
    ("τ", "Kendall rank correlation coefficient"),
    ("ρ", "Spearman rank correlation coefficient"),
]


def build(m, doc):
    """m is the make_report module, used for its paragraph helpers."""
    P, J, F = m.p, m.justified, m.field
    T, TU = m.TITLE, m.TITLE_UPPER

    # ---------------------------------------------------------- title page
    P(doc, "B.Tech. BCSE497J – Project-I", size=14, bold=True,
      align=CENTER, space_after=20)
    P(doc, TU, size=18, bold=True, align=CENTER, space_after=24,
      line_spacing=1.5)
    P(doc, "Submitted in partial fulfillment of the requirements for the "
           "degree of", align=CENTER, space_after=10)
    P(doc, "Bachelor of Technology", size=14, bold=True, align=CENTER,
      space_after=8)
    P(doc, "in", align=CENTER, space_after=8)
    P(doc, "Computer Science and Engineering", size=14, bold=True,
      align=CENTER, space_after=14)
    P(doc, "by", align=CENTER, space_after=10)

    t = doc.add_table(rows=3, cols=2)
    t.alignment = CENTER
    for row in t.rows:
        for cell, text in zip(row.cells,
                              ("<< Reg. No. >>", "<< STUDENT NAME >>")):
            par = cell.paragraphs[0]
            par.alignment = CENTER
            run = par.add_run(text)
            run.font.name, run.font.size, run.bold = m.BODY_FONT, Pt(12), True
    P(doc, "(names of team members sorted on register number)", size=10,
      italic=True, align=CENTER, space_after=20)

    P(doc, "Under the Supervision of", size=13, bold=True, align=CENTER,
      space_after=8)
    for line, bold in (("<< Project Guide Name >>", True),
                       ("<< Designation >>", False),
                       ("School of Computer Science and Engineering (SCOPE)",
                        False),
                       ("Vellore Institute of Technology, Vellore", False)):
        P(doc, line, bold=bold, align=CENTER, space_after=4)
    P(doc, "", space_after=20)
    P(doc, "October 2026", size=13, bold=True, align=CENTER)

    # ---------------------------------------------------------- declaration
    doc.add_page_break()
    P(doc, "DECLARATION", size=14, bold=True, align=CENTER, space_after=18)
    J(doc, "I hereby declare that the project entitled “" + T + "” "
           "submitted by me, for the award of the degree of Bachelor of "
           "Technology in Computer Science and Engineering to Vellore "
           "Institute of Technology, Vellore, is a record of bonafide work "
           "carried out by me under the supervision of "
           "<< Project Guide Name >>.")
    J(doc, "I further declare that the work reported in this project has not "
           "been submitted and will not be submitted, either in part or in "
           "full, for the award of any other degree or diploma in this "
           "institute or any other institute or university.")
    P(doc, "", space_after=26)
    P(doc, "Place\t: Vellore", space_after=4)
    P(doc, "Date\t: ______________", space_after=40)
    P(doc, "Signature of the Candidate", align=RIGHT)

    # ---------------------------------------------------------- certificate
    doc.add_page_break()
    P(doc, "CERTIFICATE", size=14, bold=True, align=CENTER, space_after=18)
    J(doc, "This is to certify that the project entitled “" + T + "” "
           "submitted by << Student Name (Reg. No.) >>, School of Computer "
           "Science and Engineering, Vellore Institute of Technology, Vellore, "
           "for the award of the degree of Bachelor of Technology in Computer "
           "Science and Engineering, is a record of bonafide work carried out "
           "by them under my supervision, as per the VIT code of academic and "
           "research ethics.")
    J(doc, "The contents of this report have not been submitted and will not "
           "be submitted either in part or in full, for the award of any "
           "other degree or diploma in this institute or any other institute "
           "or university. The project fulfills the requirements and "
           "regulations of the University and in my opinion meets the "
           "necessary standards for submission.")
    P(doc, "", space_after=26)
    P(doc, "Place\t: Vellore", space_after=4)
    P(doc, "Date\t: ______________", space_after=40)
    P(doc, "Signature of the Guide", align=RIGHT, space_after=40)
    P(doc, "Internal Examiner\t\t\t\t\tExternal Examiner", space_after=40)
    P(doc, "<< Head of the Department >>", align=CENTER, bold=True,
      space_after=2)
    P(doc, "Computer Science and Engineering", align=CENTER)

    # ------------------------------------------------------ acknowledgements
    doc.add_page_break()
    P(doc, "ACKNOWLEDGEMENTS", size=14, bold=True, align=CENTER,
      space_after=18)
    for para in [
        "I am deeply grateful to the management of Vellore Institute of "
        "Technology for providing me with the opportunity and the resources "
        "to undertake this project.",
        "My sincere thanks to Dr. Jaisankar N, the Dean of the School of "
        "Computer Science and Engineering, for his unwavering support and "
        "encouragement throughout the course of this work.",
        "I express my profound appreciation to << Head of the Department >>, "
        "the Head of the Department, for the insightful guidance and the "
        "continuous support extended to the project team.",
        "I am immensely thankful to my project guide, "
        "<< Project Guide Name >>, for the dedicated mentorship and the "
        "invaluable feedback that shaped this work. The patience and the "
        "questions asked at every review are the reason several comfortable "
        "assumptions in this project were measured rather than believed.",
        "Finally, I thank my family and my friends for their steady "
        "encouragement during the course of this project.",
    ]:
        J(doc, para)
    P(doc, "", space_after=32)
    P(doc, "<< Name of the Candidate >>", align=RIGHT, bold=True)

    # ------------------------------------------------------------- contents
    doc.add_page_break()
    P(doc, "TABLE OF CONTENTS", size=14, bold=True, align=CENTER,
      space_after=14)
    F(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')

    doc.add_page_break()
    P(doc, "LIST OF FIGURES", size=14, bold=True, align=CENTER, space_after=14)
    F(doc.add_paragraph(), 'TOC \\h \\z \\c "Figure"')

    doc.add_page_break()
    P(doc, "LIST OF TABLES", size=14, bold=True, align=CENTER, space_after=14)
    F(doc.add_paragraph(), 'TOC \\h \\z \\c "Table"')

    # -------------------------------------------------------- abbreviations
    doc.add_page_break()
    P(doc, "LIST OF ABBREVIATIONS", size=14, bold=True, align=CENTER,
      space_after=14)
    for short, long in ABBREVIATIONS:
        par = P(doc, space_after=2)
        r = par.add_run(short + "\t")
        r.font.name, r.font.size, r.bold = m.BODY_FONT, Pt(12), True
        r = par.add_run(long)
        r.font.name, r.font.size = m.BODY_FONT, Pt(12)

    doc.add_page_break()
    P(doc, "SYMBOLS AND NOTATIONS", size=14, bold=True, align=CENTER,
      space_after=14)
    for sym, mean in SYMBOLS:
        par = P(doc, space_after=2)
        r = par.add_run(sym + "\t")
        r.font.name, r.font.size, r.bold = m.BODY_FONT, Pt(12), True
        r = par.add_run(mean)
        r.font.name, r.font.size = m.BODY_FONT, Pt(12)
